"""Build workbooks in the audited RBNZ series layout.

The layout -- ``Data`` with group, series, notes, ``Unit`` and ``Series Id``
rows, ``Series Definitions`` and ``Table Description`` with ``Published Date``
-- is the one in the B2 daily-close workbook RBNZ published on 2026-09-15. The
OCR values used by default are the published ones in that workbook
(0.25 on 2020-05-15; 2.75 on 2026-09-08 .. 2026-09-14).
"""

from __future__ import annotations

import io
from datetime import date

import openpyxl

OCR_PUBLISHED = {
    date(2020, 5, 15): 0.25,
    date(2026, 9, 8): 2.75,
    date(2026, 9, 9): 2.75,
    date(2026, 9, 10): 2.75,
    date(2026, 9, 11): 2.75,
    date(2026, 9, 14): 2.75,
}


def workbook(
    table: str,
    columns: list[tuple[str, str, str, str]],
    rows: list[tuple[date, list[float | None]]],
    published: date = date(2026, 9, 15),
) -> bytes:
    """columns: (group, series, unit, series_id)."""
    book = openpyxl.Workbook()
    data = book.active
    assert data is not None
    data.title = "Data"
    data.append([None, *[c[0] for c in columns]])
    data.append([None, *[c[1] for c in columns]])
    data.append(["Notes", *[None] * len(columns)])
    data.append(["Unit", *[c[2] for c in columns]])
    data.append(["Series Id", *[c[3] for c in columns]])
    for when, values in rows:
        data.append([when, *values])
    definitions = book.create_sheet("Series Definitions")
    definitions.append(["Group", "Series", "Series Id", "Unit", "Note"])
    for group, series, unit, series_id in columns:
        definitions.append([group, series, series_id, unit, None])
    description = book.create_sheet("Table Description")
    description.append(["Published By", "Reserve Bank of New Zealand"])
    description.append(["Table", f"Daily table - {table}"])
    description.append(["Published Date", published])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def b2(values: dict[date, float] | None = None, published: date = date(2026, 9, 15)) -> bytes:
    """OCR between two decoys, as in the published workbook."""
    points = OCR_PUBLISHED if values is None else values
    columns = [
        ("Bank bill yields", "90 days", "%pa", "INM.DB03.NZZV"),
        ("Cash rate", "Official Cash Rate (OCR)", "%pa", "INM.DP1.N"),
        ("Swap rates close", "2 year", "%pa", "INM.DS02.NZZC"),
    ]
    rows: list[tuple[date, list[float | None]]] = [(when, [3.0, value, None]) for when, value in sorted(points.items())]
    return workbook("B2", columns, rows, published)


def b1(values: dict[date, float], unit: str = "Index", published: date = date(2026, 9, 15)) -> bytes:
    """TWI next to a decoy column. ``TEST.DECOY`` is deliberately not an RBNZ
    identifier, and the ``Index`` unit string is a test value: the live B1 unit
    text has not been observed from this environment (see METHODOLOGY.md)."""
    columns = [
        ("Exchange rates (quoted per NZ$)", "Decoy column", "USD", "TEST.DECOY"),
        ("Trade Weighted Index", "TWI (17 currencies)", unit, "EXRT.DS41.NZB17"),
    ]
    rows: list[tuple[date, list[float | None]]] = [(when, [0.6, value]) for when, value in sorted(values.items())]
    return workbook("B1", columns, rows, published)
